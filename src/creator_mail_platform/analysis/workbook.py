from __future__ import annotations

from datetime import UTC, datetime
from pathlib import Path
from typing import Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill("solid", fgColor="273C75")
HEADER_FONT = Font(color="FFFFFF", bold=True)
ALT_FILL = PatternFill("solid", fgColor="F2F5FA")
HYPERLINK_FONT = Font(color="0563C1", underline="single")
DATE_FORMAT = "yyyy-mm-dd hh:mm:ss"


def write_active_creator_workbook(
    output_path: Path,
    *,
    active_rows: list[dict[str, object]],
    reply_rows: list[dict[str, object]],
) -> None:
    """Replace the generated workbook atomically; it contains no manual columns."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    default = workbook.active
    workbook.remove(default)
    _write_sheet(workbook, "活跃达人", active_rows)
    _write_sheet(workbook, "回信分析", reply_rows)

    temporary = output_path.with_name(f"{output_path.stem}.tmp{output_path.suffix}")
    workbook.save(temporary)
    # Reopen once so broken output is never promoted to the business workbook.
    checked = load_workbook(temporary, read_only=True, data_only=False)
    try:
        if checked.sheetnames != ["活跃达人", "回信分析"]:
            raise RuntimeError("Excel sheet verification failed")
    finally:
        checked.close()
    temporary.replace(output_path)


def _write_sheet(workbook: Workbook, title: str, rows: list[dict[str, object]]) -> None:
    sheet = workbook.create_sheet(title)
    headers = list(rows[0]) if rows else _empty_headers(title)
    sheet.append(headers)
    for row_index, row in enumerate(rows, start=2):
        sheet.append([_excel_value(row.get(header)) for header in headers])
        if row_index % 2 == 0:
            for cell in sheet[row_index]:
                cell.fill = ALT_FILL
        if "profile_url" in headers:
            url_cell = sheet.cell(row=row_index, column=headers.index("profile_url") + 1)
            if isinstance(url_cell.value, str) and url_cell.value.startswith(
                ("https://", "http://")
            ):
                url_cell.hyperlink = url_cell.value
                url_cell.font = HYPERLINK_FONT

    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 30

    for column_index, header in enumerate(headers, start=1):
        values = [str(header)] + [str(row.get(header) or "") for row in rows[:200]]
        width = min(max(max(len(value) for value in values) + 2, 12), _column_cap(str(header)))
        sheet.column_dimensions[get_column_letter(column_index)].width = width
        for cell in sheet[get_column_letter(column_index)][1:]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(cell.value, datetime):
                cell.number_format = DATE_FORMAT


def _column_cap(header: str) -> int:
    if header in {"profile_bio", "analysis_note", "sent_body_text", "original_reply_text"}:
        return 60
    if header in {"additional_info", "email_feedback", "reply_summary_zh"}:
        return 45
    return 32


def _excel_value(value: object) -> object:
    if isinstance(value, datetime):
        if value.tzinfo is not None:
            return value.astimezone(UTC).replace(tzinfo=None)
        return value
    return value


def _empty_headers(title: str) -> Sequence[str]:
    if title == "活跃达人":
        return ("creator_id", "platform", "handle", "activation_source")
    return ("creator_id", "platform", "handle", "intent")
